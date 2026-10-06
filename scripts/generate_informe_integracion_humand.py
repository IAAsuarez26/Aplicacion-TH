# -*- coding: utf-8 -*-
"""
Generador del Informe Gerencial: Proceso de Integración y Sincronización TH <-> Humand
Formato: Microsoft Word (.docx) con tipografía ejecutiva, tablas formateadas, callouts y métricas.
Versión: 2.5 (Actualizada con Sincronización Selectiva por Campos y Fechas de Nacimiento)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Paleta Corporativa Ponce & Benzo
COLOR_NAVY_HEX = "1B365D"       # Azul Marino Principal
COLOR_BLUE_HEX = "2B6CB0"       # Azul Corporativo Medio
COLOR_ACCENT_HEX = "D97706"     # Ámbar Dorado Ejecutivo
COLOR_GREEN_HEX = "2F855A"      # Verde Éxito / Sincronizado
COLOR_RED_HEX = "C53030"        # Rojo Alerta / Restricción
COLOR_BG_LIGHT_HEX = "F8FAFC"   # Fondo tarjeta / callout
COLOR_BG_HEADER_HEX = "EDF2F7"  # Fondo encabezado tabla
COLOR_BORDER_HEX = "CBD5E1"     # Borde sutil slate

COLOR_NAVY = RGBColor(27, 54, 93)
COLOR_BLUE = RGBColor(43, 108, 176)
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)
COLOR_WHITE = RGBColor(255, 255, 255)

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_title(doc, text, subtitle_text):
    # Encabezado institucional previo
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_meta = p_meta.add_run("PONCE & BENZO | DIRECCIÓN DE SISTEMAS Y TECNOLOGÍA\nGERENCIA CORPORATIVA DE TALENTO HUMANO")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(8.5)
    run_meta.font.bold = True
    run_meta.font.color.rgb = COLOR_TEXT_MUTED

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run(subtitle_text)
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = COLOR_BLUE

    # Línea decorativa
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(16)
    pPr = p_line._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="2" w:color="{COLOR_NAVY_HEX}"/></w:pBdr>')
    pPr.append(pBdr)

def add_heading_styled(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.color.rgb = COLOR_NAVY
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="10" w:space="4" w:color="{COLOR_BLUE_HEX}"/></w:pBdr>')
        pPr.append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_BLUE
    elif level == 3:
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_TEXT_MAIN

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = COLOR_TEXT_MAIN
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT_MAIN
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = COLOR_TEXT_MAIN
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT_MAIN

def add_callout(doc, title, text, border_color_hex=COLOR_BLUE_HEX, bg_color_hex=COLOR_BG_LIGHT_HEX):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tblPr = tbl._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color_hex}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
        <w:insideH w:val="none"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r_title = p.add_run(title)
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = COLOR_NAVY

    p_body = cell.add_paragraph()
    p_body.paragraph_format.space_after = Pt(0)
    p_body.paragraph_format.line_spacing = 1.15
    r_text = p_body.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_TEXT_MAIN

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)

def style_table_header(row, col_widths=None):
    for i, cell in enumerate(row.cells):
        set_cell_background(cell, COLOR_NAVY_HEX)
        set_cell_margins(cell, top=130, bottom=130, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = COLOR_WHITE

def style_table_rows(table, col_widths=None, alternate_bg=True):
    for r_idx, row in enumerate(table.rows[1:]):
        bg = COLOR_BG_LIGHT_HEX if (alternate_bg and r_idx % 2 == 1) else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)
                    r.font.color.rgb = COLOR_TEXT_MAIN

def generate_document():
    doc = docx.Document()

    # Configuración de Márgenes (Carta estándar)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ---------------------------------------------------------
    # ENCABEZADO Y TÍTULO
    # ---------------------------------------------------------
    add_title(
        doc,
        "Informe Gerencial de Arquitectura e Integración TH ↔ Humand",
        "Modelo de Sincronización Selectiva por Campos, Carga de Fechas de Nacimiento, Premisas Técnicas y Gobierno de Datos"
    )

    # Tabla Ficha Técnica
    ficha_table = doc.add_table(rows=4, cols=2)
    ficha_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(2.2), Inches(4.3)]
    
    meta_data = [
        ("Proyecto / Iniciativa:", "Plataforma Integral de Talento Humano — Ecosistema Digital Ponce & Benzo"),
        ("Destinatarios:", "Dirección General, Dirección de Sistemas & Tecnología, Gerencia de Talento Humano"),
        ("Fecha de Emisión:", "Octubre 2026 | Versión Oficial 2.5 (Actualizada)"),
        ("Clasificación de Seguridad:", "Confidencial — Uso Interno Corporativo"),
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = ficha_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = col_widths[0], col_widths[1]
        set_cell_background(cell_lbl, COLOR_BG_HEADER_HEX)
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, top=70, bottom=70, left=120, right=120)
        set_cell_margins(cell_val, top=70, bottom=70, left=120, right=120)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_NAVY
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        
    set_table_borders(ficha_table, color="CBD5E1", sz="4", val="single")
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ---------------------------------------------------------
    # RESUMEN EJECUTIVO
    # ---------------------------------------------------------
    add_heading_styled(doc, "Resumen Ejecutivo", level=1)
    add_body_paragraph(
        doc,
        "El presente informe establece con rigor técnico y funcional la arquitectura de integración y el modelo de sincronización "
        "entre la base de datos central de la Aplicación de Talento Humano (Aplicación TH) y la plataforma de comunicación y experiencia del empleado "
        "Humand. El objetivo medular es garantizar un intercambio de información transparente, seguro, automatizado y bajo estricto control corporativo."
    )
    add_body_paragraph(
        doc,
        "En esta versión actualizada 2.5, se documentan formalmente dos hitos trascendentales en la madurez del ecosistema: "
        "(1) La integración del Módulo de Sincronización Selectiva por Campos (Delta PATCH), que permite a la Gerencia seleccionar libremente "
        "mediante checkboxes interactivos qué atributos específicos desea actualizar en Humand (por ejemplo: únicamente la fecha de nacimiento, el correo o el teléfono), "
        "sin alterar departamentos, puestos, credenciales ni el resto del perfil; y (2) La digitalización, conciliación y carga del 100% de las Fechas de Nacimiento "
        "en la base de datos de TH para los 187 colaboradores activos, homologándolas con el atributo nativo 'birthdate' en Humand."
    )
    add_body_paragraph(
        doc,
        "Como principio rector, se ratifica que el sistema NO está obligado a realizar sincronizaciones totales masivas. La arquitectura ofrece "
        "capacidades duales de máxima flexibilidad: permite cargas globales o estructurales en cascada cuando se crean departamentos o cargos masivos, "
        "pero prioriza la Sincronización Delta Quirúrgica para la operación habitual, protegiendo las cuotas de Rate Limit de la API y garantizando "
        "la total inocuidad sobre los expedientes existentes en la nube."
    )

    # ---------------------------------------------------------
    # SECCIÓN 1: SINCRONIZACIÓN TOTAL VS. PARCIAL SELECTIVA
    # ---------------------------------------------------------
    add_heading_styled(doc, "1. Modelo de Sincronización: Cascada Completa vs. Sincronización Selectiva", level=1)
    add_body_paragraph(
        doc,
        "Una de las premisas arquitectónicas más críticas en la integración de sistemas de recursos humanos es la eficiencia, la precisión y el control de tráfico en las APIs. "
        "En la integración TH ↔ Humand, los administradores cuentan con la libertad de elegir entre modalidades operativas diferenciadas según la necesidad:"
    )

    add_heading_styled(doc, "1.1 Sincronización Selectiva por Campos (Nueva Funcionalidad Delta PATCH)", level=2)
    add_body_paragraph(
        doc,
        "La Sincronización Selectiva por Campos permite al usuario escoger exactamente cuáles atributos de datos desea transmitir hacia Humand. "
        "Esta modalidad opera bajo el verbo HTTP PATCH, lo que significa que la API de Humand únicamente actualiza los campos enviados en el payload y "
        "deja intacto el 100% de los datos preexistentes (cargos, departamentos, contraseñas, antigüedad, organigrama, etc.)."
    )
    add_bullet(doc, "Disponibles en el nuevo diálogo interactivo de la interfaz: (a) Solo Fecha de Nacimiento; (b) Datos Personales y Biográficos (Nacimiento, Teléfono); (c) Datos Laborales (Fecha de Ingreso, Correo, Supervisor BOSS); (d) Todos los Campos; o (e) Selección Personalizada libre.", bold_prefix="Presets Operativos de Selección Rápida: ")
    add_bullet(doc, "Si el usuario selecciona solo 'Fecha de Nacimiento', el sistema envía únicamente { 'birthdate': 'AAAA-MM-DD' }. No se tocan correos, teléfonos, roles ni membresías organizacionales.", bold_prefix="Inocuidad y Protección de Datos: ")
    add_bullet(doc, "Ahorra drásticamente el consumo de peticiones por minuto (Rate Limit) al evitar llamadas innecesarias a endpoints de departamentos y cargos.", bold_prefix="Alta Eficiencia en API: ")

    add_heading_styled(doc, "1.2 Sincronización Parcial por Colaborador Individual", level=2)
    add_body_paragraph(
        doc,
        "Tanto en la tabla principal de colaboradores como en el modal de sincronización, es posible disparar la actualización individual para un solo empleado "
        "(utilizando su cédula de identidad como identificador unívoco employeeInternalId). Esto resulta ideal para nuevos ingresos o correcciones puntuales, "
        "evitando procesar a toda la nómina para un solo caso."
    )

    add_heading_styled(doc, "1.3 Sincronización en Cascada Completa (Estructural)", level=2)
    add_body_paragraph(
        doc,
        "La Sincronización en Cascada (PUT /users + asignación de membresías) se reserva para momentos en los que se han creado nuevos departamentos o puestos en TH "
        "que deben aprovisionarse primero en Humand, o durante auditorías generales mensuales de consistencia referencial."
    )

    # Tabla Comparativa Total vs Parcial
    add_heading_styled(doc, "1.4 Matriz Comparativa de Métodos de Sincronización", level=2)
    comp_table = doc.add_table(rows=6, cols=4)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.6), Inches(1.3), Inches(1.8), Inches(1.8)]
    
    headers = ["Modalidad", "Endpoint / Verbo", "Alcance de Datos", "Consumo de API / Cuota"]
    for i, h in enumerate(headers):
        comp_table.rows[0].cells[i].paragraphs[0].text = h
    style_table_header(comp_table.rows[0], widths)

    rows_data = [
        ("Sincronización Selectiva por Campos (NUEVO)", "PATCH /users/{id}", "Modifica exclusivamente los campos elegidos por el usuario (ej. solo fecha de nacimiento).", "Mínimo; 1 sola llamada ligera por usuario sin tocar otros datos."),
        ("Actualización de Campos de Perfil", "PATCH /users/{id}/profile-fields", "Afecta únicamente los campos dinámicos (Estado Civil, Nivel Educativo).", "Ultra ligero; preserva exactamente el resto del perfil intacto."),
        ("Asignación de Segmentaciones", "POST /segmentations/users", "Asigna grupos de Ubicación, Género o Sede geográfica.", "Específico; requiere que el grupo exista previamente en Humand."),
        ("Asignación de Estructura Organizacional", "PUT /departments/members\nPUT /job-positions/members", "Enlaza al empleado con su departamento o puesto de trabajo correspondiente.", "Moderado; asegura que los organigramas reflejen la verdad de TH."),
        ("Sincronización en Cascada Completa", "PUT /users + Membresías", "Reemplazo estructural del usuario y reasignación de dependencias.", "Alto tráfico; reservado para aprovisionamiento inicial o reestructuraciones.")
    ]

    for idx, r_data in enumerate(rows_data):
        row = comp_table.rows[idx + 1]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].paragraphs[0].text = val
    style_table_rows(comp_table, widths)

    add_callout(
        doc,
        "Conclusión Operativa:",
        "La Gerencia de Talento Humano NO está obligada a realizar sincronizaciones globales ni pesadas. "
        "La nueva funcionalidad de Sincronización Selectiva por Campos permite enviar puntualmente la Fecha de Nacimiento (o cualquier combinación de campos) "
        "con total tranquilidad, garantizando 100% de inocuidad sobre la estructura organizativa existente.",
        border_color_hex=COLOR_GREEN_HEX
    )

    # ---------------------------------------------------------
    # SECCIÓN 2: CARACTERÍSTICAS Y PILARES DE LA INTEGRACIÓN
    # ---------------------------------------------------------
    add_heading_styled(doc, "2. Características de la Integración", level=1)
    add_body_paragraph(
        doc,
        "La integración entre Aplicación TH y Humand Public API v1 se fundamenta en los siguientes pilares de arquitectura de software y gestión empresarial:"
    )

    add_bullet(doc, "La Aplicación TH (PostgreSQL en InsForge) actúa como la única fuente autoritativa de datos. Humand es un sistema satélite de experiencia, comunicación y autoservicio. Ningún dato manual creado en Humand sobrescribe la base central de TH.", bold_prefix="1. Unidireccionalidad de Autoridad (SSOT): ")
    add_bullet(doc, "La comunicación se realiza mediante HTTPS REST API contra https://api-prod.humand.co/public/api/v1 utilizando autenticación estricta en cabecera 'Authorization: Basic {HUMAND_API_KEY}'.", bold_prefix="2. Protocolo Seguro y Autenticación: ")
    add_bullet(doc, "La cédula de identidad del colaborador (sanitizada numéricamente sin letras V/E, puntos ni guiones, ej. '14714667') actúa como el employeeInternalId en Humand. Esta es la llave unívoca e inmutable que indexa todas las operaciones.", bold_prefix="3. Identificador Único Universal: ")
    add_bullet(doc, "Humand organiza los perfiles mediante tres capas complementarias: (a) Atributos nativos del core (nombre, correo, teléfono, fecha de contratación, fecha de nacimiento); (b) Segmentaciones grupales jerárquicas (Ubicación, Género, Sede); y (c) Profile Fields dinámicos asignados por UUIDs globales.", bold_prefix="4. Soporte Multidimensional de Datos: ")
    add_bullet(doc, "Al sincronizar la línea de mando (di_supervisor y di_evaluador), Humand construye el organigrama empresarial de manera automática mediante la asignación de relaciones BOSS y REVIEWER, generando la subordinación bidireccional correspondiente.", bold_prefix="5. Organigrama Dinámico Automatizado: ")

    # ---------------------------------------------------------
    # SECCIÓN 3: PREMISAS FUNDAMENTALES DE FUNCIONAMIENTO
    # ---------------------------------------------------------
    add_heading_styled(doc, "3. Premisas Fundamentales de Funcionamiento", level=1)
    add_body_paragraph(
        doc,
        "Para garantizar la integridad y estabilidad de la nómina y de la plataforma social corporativa, la integración opera bajo reglas inviolables:"
    )

    add_heading_styled(doc, "3.1 Gobierno Interno con el Campo estatus_h (Regla de Exclusión)", level=2)
    add_body_paragraph(
        doc,
        "Por solicitud de la Gerencia, se ha desacoplado la condición de sincronización respecto al estado laboral convencional (ACTIVO/INACTIVO). "
        "Se implementó el campo estatus_h en las tablas maestras de Empleados, Departamentos y Cargos bajo las siguientes premisas:"
    )
    add_bullet(doc, "El registro está validado y autorizado para transmitirse hacia Humand durante cualquier proceso de sincronización.", bold_prefix="estatus_h = 1 (Habilitado Humand): ")
    add_bullet(doc, "El registro pertenece única y exclusivamente a la Aplicación TH (nóminas confidenciales, personal externo, directores no vinculados o registros en borrador). NUNCA viajará ni existirá en Humand.", bold_prefix="estatus_h = 0 (Solo TH / Excluido): ")
    add_bullet(doc, "El campo estatus_h NO existe en Humand, no se crea como profile field, ni se muestra a los usuarios en la aplicación móvil ni web. Funciona estrictamente tras bastidores como una compuerta lógica interna.", bold_prefix="Invisibilidad Absoluta en Humand: ")

    add_callout(
        doc,
        "Auditoría Actual de la Nómina (Octubre 2026):",
        "Total Colaboradores en Aplicación TH: 187 colaboradores activos.\n"
        "• Cobertura de Fechas de Nacimiento en BD: 187 de 187 colaboradores (100.0% de cobertura confirmada).\n"
        "• Colaboradores asignados y vinculados en Humand: 163 personas principales (100% coincidentes con Humand).\n"
        "• Colaboradores en fase de homologación de correo corporativo: 13 colaboradores pendientes.\n"
        "• Colaboradores exclusivos o confidenciales de TH: Protegidos bajo estatus_h = 0.",
        border_color_hex=COLOR_NAVY_HEX
    )

    add_heading_styled(doc, "3.2 Política Antispam en Correo Electrónico", level=2)
    add_body_paragraph(
        doc,
        "La creación de usuarios en Humand mediante el endpoint POST /users dispara automáticamente un correo electrónico de invitación al empleado si posee dirección de correo. "
        "Para evitar saturar a los colaboradores o enviar correos en fases de prueba, la integración utiliza PUT /users o PATCH /users/{id}. "
        "Bajo estos métodos, los usuarios se aprovisionan o actualizan de forma completamente silenciosa, permitiendo que la Gerencia de Talento Humano "
        "decida el momento exacto para lanzar formalmente la invitación corporativa o capacitaciones."
    )

    add_heading_styled(doc, "3.3 Protección Incondicional de Cuentas Técnicas", level=2)
    add_body_paragraph(
        doc,
        "Los motores de conciliación y sincronización tienen programada una lista blanca (Whitelisting) inalterable. Bajo ninguna circunstancia se modifican, desactivan "
        "ni eliminan las cuentas técnicas y de administración de la plataforma:"
    )
    add_bullet(doc, "Cuenta de acompañamiento técnico y consultoría de implementación de Humand.", bold_prefix="carolina.ellena@humand.co: ")
    add_bullet(doc, "Usuario del sistema y servicio asociado a la API Key con permisos de lectura/escritura.", bold_prefix="integracionespb: ")
    add_bullet(doc, "Cuenta oficial para la publicación de noticias, reconocimientos y banners institucionales.", bold_prefix="comunicacionespb: ")

    add_heading_styled(doc, "3.4 Resiliencia ante Rate Limits de la API", level=2)
    add_body_paragraph(
        doc,
        "Humand impone un límite de seguridad de 50 peticiones por minuto (RateLimit-Policy: 50;w=60). Si este límite se sobrepasa, la API responde con HTTP 429 Too Many Requests. "
        "El motor de sincronización de la Aplicación TH incorpora automáticamente lógica de reintentos con Backoff Exponencial y lectura de la cabecera 'Retry-After', "
        "además de una pausa programada de 200 ms entre colaboradores, garantizando que ninguna transacción se pierda y que el proceso se ejecute con total fluidez."
    )

    # ---------------------------------------------------------
    # SECCIÓN 4: MAPEO DE CAMPOS NATIVOS Y EXTENDIDOS
    # ---------------------------------------------------------
    add_heading_styled(doc, "4. Catálogo y Mapeo de Campos Nativos y Extendidos", level=1)
    add_body_paragraph(
        doc,
        "Para dar cumplimiento integral a los requerimientos de la Gerencia, se implementaron en la base de datos de TH y se sincronizan con Humand "
        "los siguientes atributos corporativos, mapeados según los estándares exactos de la API:"
    )

    fields_table = doc.add_table(rows=5, cols=4)
    fields_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_widths = [Inches(1.5), Inches(1.5), Inches(1.8), Inches(1.7)]
    
    f_headers = ["Campo en TH", "Destino en Humand", "Identificador Técnico Humand", "Valores / Formato Homologado"]
    for i, h in enumerate(f_headers):
        fields_table.rows[0].cells[i].paragraphs[0].text = h
    style_table_header(fields_table.rows[0], f_widths)

    fields_data = [
        ("Fecha de Nacimiento\n(fecha_nacimiento)\n[NUEVO]", "Atributo Nativo Core\n('birthdate')", "Parámetro direct: birthdate\n(PATCH /users/{id})", "Formato ISO 'AAAA-MM-DD'\n(ej. '1995-04-19'). 100% poblado en BD."),
        ("Ubicación\n(ubicacion)", "Segmentación Grupal\n('Ubicación')", "Group ID: 451900\n(POST /segmentations/users)", "Caracas, Yagua, Barquisimeto, Maracaibo, Valencia, Puerto Ordaz, San Cristóbal"),
        ("Estado Civil\n(edo_civil)", "Profile Field\n(Confidencial)", "UUID:\n52005932-0bdb-438d-834c-82d8e3330e26", "Soltero, Casado, Divorciado, Viudo, Concubinato"),
        ("Nivel Educativo\n(nivel_educativo)", "Profile Field\n(Visible/Público)", "UUID:\na0b1e6b5-4bc7-444b-a27c-2844dd5a9532", "Bachiller, Técnico Medio, Técnico Superior, Universitario, Posgrado")
    ]

    for idx, r_data in enumerate(fields_data):
        row = fields_table.rows[idx + 1]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].paragraphs[0].text = val
    style_table_rows(fields_table, f_widths)

    add_body_paragraph(
        doc,
        "Nota Técnica de Homologación y Exportación: En la Aplicación TH, los selectores de la interfaz gráfica y los esquemas de base de datos "
        "fueron restringidos para coincidir al 100% con los estándares de Humand. Adicionalmente, el maestro de empleados ahora incluye la "
        "Fecha de Nacimiento y la Fecha de Ingreso en columnas dedicadas independientes, y segrega el Estado Civil y Nivel Educativo en campos separados "
        "en la exportación CSV para facilitar reportes gerenciales y de nómina."
    )

    # ---------------------------------------------------------
    # SECCIÓN 5: CICLO DE VIDA DEL COLABORADOR
    # ---------------------------------------------------------
    add_heading_styled(doc, "5. Ciclo de Vida del Colaborador en la Integración", level=1)
    add_body_paragraph(
        doc,
        "La integración acompaña todas las etapas de la relación laboral del colaborador, actuando según el evento que ocurra en la Aplicación TH:"
    )

    add_bullet(doc, "Al registrar un nuevo empleado en TH con estatus_h = 1, el motor lo inserta en Humand vía PUT /users sin disparar correos automáticos. Inmediatamente se le asocia su Ubicación geográfica, su Estado Civil, Nivel Educativo, Fecha de Nacimiento, Departamento y Puesto de Trabajo.", bold_prefix="A. Alta o Contratación: ")
    add_bullet(doc, "Cuando se edita un campo puntual (por ejemplo: carga de fecha de nacimiento, nuevo teléfono o cambio de correo), el sistema ejecuta una llamada PATCH selectiva. Solo viaja el dato modificado, manteniendo intacto el expediente histórico.", bold_prefix="B. Modificación o Movimiento Interno: ")
    add_bullet(doc, "Si un colaborador finaliza su relación laboral y se marca INACTIVO en TH, la integración ejecuta el endpoint POST /users/{id}/deactivate especificando el motivo (Renuncia, Fin de Contrato, Jubilación). Fundamentalmente, permite transferir la jefatura de los subordinados a un nuevo líder mediante el parámetro newBossId, evitando que los equipos queden sin supervisor en el organigrama.", bold_prefix="C. Baja o Desvinculación Formal (Offboarding): ")

    # ---------------------------------------------------------
    # SECCIÓN 6: CONCLUSIONES Y RECOMENDACIONES
    # ---------------------------------------------------------
    add_heading_styled(doc, "6. Conclusiones y Próximos Pasos Estratégicos", level=1)
    add_body_paragraph(
        doc,
        "1. Flexibilidad y Autonomía Absoluta: La Dirección de Talento Humano cuenta ahora con una herramienta de Sincronización Selectiva por Campos "
        "que le permite actualizar en Humand exactamente lo que necesita, cuando lo necesita, sin depender de procesos masivos pesados ni alterar la estructura organizativa."
    )
    add_body_paragraph(
        doc,
        "2. Cobertura Total de Fechas de Nacimiento: Se ha cerrado la brecha de información biográfica en la nómina. Con el 100% de las fechas cargadas en TH "
        "y el endpoint PATCH configurado, el ecosistema está listo para celebrar cumpleaños corporativos y activar dinámicas de fidelización en Humand."
    )
    add_body_paragraph(
        doc,
        "3. Plataforma Lista para Fases Siguientes: Al tener los catálogos, puestos, departamentos, segmentaciones de ubicación y datos extendidos 100% alineados, "
        "el ecosistema queda completamente preparado para conectar los módulos avanzados de Humand:"
    )
    add_bullet(doc, "Para la consulta de solicitudes de vacaciones aprobadas y su descuento automático en el balance de días de TH.", bold_prefix="• Módulo Time-Off: ")
    add_bullet(doc, "Para la lectura programática de marcas de entrada y salida orientadas a la prenómina.", bold_prefix="• Módulo Time-Tracking: ")
    add_bullet(doc, "Para la distribución digital y segura de recibos de pago en PDF directamente al perfil móvil de cada empleado.", bold_prefix="• Módulo Documentos de Nómina: ")

    # Firma de Responsabilidad Técnica
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.line_spacing = 1.15
    r_s1 = p_sign.add_run("DIRECCIÓN DE SISTEMAS Y TECNOLOGÍA\n")
    r_s1.bold = True
    r_s1.font.name = "Calibri"
    r_s1.font.size = Pt(10)
    r_s1.font.color.rgb = COLOR_NAVY
    
    r_s2 = p_sign.add_run("Liderazgo de Arquitectura Cloud e Integraciones Empresariales\nCorporación Ponce & Benzo")
    r_s2.font.name = "Calibri"
    r_s2.font.size = Pt(9.5)
    r_s2.font.color.rgb = COLOR_TEXT_MUTED

    # Guardar documento
    output_path = os.path.join(os.path.dirname(__file__), "..", "documentos", "Informe Gerencial - Proceso de Integracion y Sincronizacion TH-Humand.docx")
    doc.save(output_path)
    print(f"Documento generado exitosamente en: {output_path}")

if __name__ == "__main__":
    generate_document()
