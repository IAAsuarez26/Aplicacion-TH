# -*- coding: utf-8 -*-
"""
Script de Auditoría y Exportación de Catálogos desde Humand hacia Excel
======================================================================
Descarga todos los Departamentos y Puestos de Trabajo directamente de la API
de Humand y genera un reporte en Excel para su auditoría y validación.
"""

import os
import json
import urllib.request
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HUMAND_BASE_URL = "https://api-prod.humand.co/public/api/v1"

def load_env():
    env_path = os.path.join(os.path.dirname(__file__), "..", ".env.local")
    if os.path.exists(env_path):
        with open(env_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    k = k.strip()
                    v = v.strip().strip('"').strip("'")
                    if k not in os.environ:
                        os.environ[k] = v

def humand_get_all(endpoint, api_key):
    items = []
    page = 1
    limit = 50
    headers = {
        "Authorization": f"Basic {api_key}",
        "Content-Type": "application/json"
    }
    while True:
        url = f"{HUMAND_BASE_URL}{endpoint}?page={page}&limit={limit}"
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=20) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            current = data.get("items", [])
            items.extend(current)
            total_pages = data.get("totalPages", 1)
            if page >= total_pages or not current:
                break
            page += 1
    return items

def create_excel_report():
    load_env()
    api_key = os.getenv("HUMAND_API_KEY")
    if not api_key:
        print("Error: HUMAND_API_KEY no encontrada.")
        return
        
    print("Consultando datos en vivo desde Humand...")
    deps = humand_get_all("/departments", api_key)
    cargos = humand_get_all("/job-positions", api_key)
    
    print(f"Recuperados {len(deps)} departamentos y {len(cargos)} puestos de trabajo.")
    
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto
    wb.remove(wb.active)
    
    # Estilos
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E0"),
        right=Side(style="thin", color="CBD5E0"),
        top=Side(style="thin", color="CBD5E0"),
        bottom=Side(style="thin", color="CBD5E0")
    )
    
    # ----------------------------------------------------
    # Hoja 1: Departamentos
    # ----------------------------------------------------
    ws_dep = wb.create_sheet(title="Departamentos (44)")
    ws_dep.views.sheetView[0].showGridLines = True
    
    # Título
    ws_dep.merge_cells("A1:E1")
    title_cell = ws_dep["A1"]
    title_cell.value = "AUDITORÍA DE DEPARTAMENTOS EN HUMAND - PONCE & BENZO"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1B365D")
    title_cell.alignment = Alignment(vertical="center")
    ws_dep.row_dimensions[1].height = 30
    
    headers_dep = ["N°", "ID Humand", "Nombre del Departamento", "Usuarios Vinculados", "Fecha Creación (UTC)"]
    ws_dep.append(headers_dep)
    ws_dep.row_dimensions[2].height = 24
    
    for col_num, h in enumerate(headers_dep, 1):
        c = ws_dep.cell(row=2, column=col_num)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        
    deps_sorted = sorted(deps, key=lambda x: x.get("name", "").lower())
    for idx, d in enumerate(deps_sorted, 1):
        row_num = idx + 2
        r_data = [
            idx,
            d.get("id"),
            d.get("name"),
            d.get("usersCount", 0),
            str(d.get("createdAt", "")).replace("T", " ")[:19]
        ]
        ws_dep.append(r_data)
        ws_dep.row_dimensions[row_num].height = 20
        
        is_even = idx % 2 == 0
        for col_num in range(1, 6):
            cell = ws_dep.cell(row=row_num, column=col_num)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if is_even:
                cell.fill = zebra_fill
            if col_num in [1, 2, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Auto ancho de columnas
    for col in ws_dep.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_dep.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # ----------------------------------------------------
    # Hoja 2: Puestos de Trabajo
    # ----------------------------------------------------
    ws_car = wb.create_sheet(title="Puestos de Trabajo (97)")
    ws_car.views.sheetView[0].showGridLines = True
    
    ws_car.merge_cells("A1:E1")
    t2 = ws_car["A1"]
    t2.value = "AUDITORÍA DE PUESTOS DE TRABAJO (CARGOS) EN HUMAND - PONCE & BENZO"
    t2.font = Font(name="Calibri", size=14, bold=True, color="1B365D")
    t2.alignment = Alignment(vertical="center")
    ws_car.row_dimensions[1].height = 30
    
    headers_car = ["N°", "ID Humand", "Nombre del Puesto / Cargo", "Usuarios Vinculados", "Fecha Creación (UTC)"]
    ws_car.append(headers_car)
    ws_car.row_dimensions[2].height = 24
    
    for col_num, h in enumerate(headers_car, 1):
        c = ws_car.cell(row=2, column=col_num)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        
    cargos_sorted = sorted(cargos, key=lambda x: x.get("name", "").lower())
    for idx, c in enumerate(cargos_sorted, 1):
        row_num = idx + 2
        r_data = [
            idx,
            c.get("id"),
            c.get("name"),
            c.get("usersCount", 0),
            str(c.get("createdAt", "")).replace("T", " ")[:19]
        ]
        ws_car.append(r_data)
        ws_car.row_dimensions[row_num].height = 20
        
        is_even = idx % 2 == 0
        for col_num in range(1, 6):
            cell = ws_car.cell(row=row_num, column=col_num)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if is_even:
                cell.fill = zebra_fill
            if col_num in [1, 2, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    for col in ws_car.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_car.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    output_path = os.path.join(os.path.dirname(__file__), "..", "documentos", "Reporte_Verificacion_Catalogos_Humand.xlsx")
    wb.save(output_path)
    print(f"Reporte Excel generado con éxito en: {output_path}")

if __name__ == "__main__":
    create_excel_report()
