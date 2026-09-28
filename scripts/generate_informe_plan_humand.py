# -*- coding: utf-8 -*-
"""
Generador del Informe Gerencial: Plan de Implementación y Pruebas de Integración TH <-> Humand
Formato: Microsoft Word (.docx) con diseño ejecutivo y confidencialidad salarial estricta.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Paleta de Colores Corporativos
COLOR_PRIMARY_HEX = "1B365D"       # Azul Marino Corporativo
COLOR_SECONDARY_HEX = "2B6CB0"     # Azul Medio
COLOR_ACCENT_HEX = "C53030"        # Rojo Alerta / DLP
COLOR_BG_LIGHT_HEX = "F7FAFC"      # Gris ultra claro
COLOR_BG_HEADER_HEX = "EDF2F7"     # Gris encabezado tabla
COLOR_BORDER_HEX = "CBD5E0"        # Borde sutil
COLOR_TEXT_MAIN_HEX = "2D3748"     # Gris oscuro texto principal
COLOR_TEXT_MUTED_HEX = "718096"    # Gris texto secundario

COLOR_PRIMARY = RGBColor(27, 54, 93)
COLOR_SECONDARY = RGBColor(43, 108, 176)
COLOR_ACCENT = RGBColor(197, 48, 48)
COLOR_TEXT_MAIN = RGBColor(45, 55, 72)
COLOR_TEXT_MUTED = RGBColor(113, 128, 150)

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_header_styled(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = COLOR_PRIMARY
        # Barra decorativa inferior
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="{COLOR_PRIMARY_HEX}"/></w:pBdr>')
        pPr.append(pBdr)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_SECONDARY
    elif level == 3:
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_TEXT_MAIN
    return p

def add_callout_box(doc, title, text, border_color="2B6CB0", bg_color="F7FAFC"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"{title}: ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor.from_string(border_color)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_TEXT_MAIN
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_document():
    doc = docx.Document()
    
    # Configuración de márgenes estándar
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    # --------------------------------------------------------------------------
    # Encabezado Superior / Membrete
    # --------------------------------------------------------------------------
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(0)
    p_pre.paragraph_format.space_after = Pt(2)
    r_pre = p_pre.add_run("DIRECCIÓN DE SISTEMAS Y TECNOLOGÍA | GERENCIA DE TALENTO HUMANO")
    r_pre.font.name = "Calibri"
    r_pre.font.size = Pt(8.5)
    r_pre.bold = True
    r_pre.font.color.rgb = COLOR_TEXT_MUTED
    
    # Título Principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("INFORME GERENCIAL: PLAN DE IMPLEMENTACIÓN Y PRUEBAS DE INTEGRACIÓN")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(17)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    
    # Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("Protocolo Canónico para la Integración de Datos TH ↔ Humand Public API v1 bajo Estricta Política de Confidencialidad Salarial")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = COLOR_SECONDARY
    
    # Tabla de Metadatos del Documento
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="CBD5E0", sz="4")
    
    meta_data = [
        ("Código de Referencia:", "INF-GER-TH-HUMAND-2026-03"),
        ("Áreas Responsables:", "Dirección de Sistemas y Tecnología / Gerencia Corporativa de Talento Humano"),
        ("Fecha de Emisión:", "28 de Septiembre de 2026"),
        ("Directiva Especial de Seguridad:", "CONFIDENCIALIDAD TOTAL: Exclusión de Bandas Salariales y Compensación")
    ]
    
    col_widths = [Inches(2.3), Inches(4.2)]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        
        c0 = row.cells[0]
        c0.width = col_widths[0]
        set_cell_background(c0, COLOR_BG_HEADER_HEX)
        set_cell_margins(c0, top=90, bottom=90, left=120, right=120)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.bold = True
        r0.font.color.rgb = COLOR_TEXT_MAIN
        
        c1 = row.cells[1]
        c1.width = col_widths[1]
        set_cell_background(c1, COLOR_BG_LIGHT_HEX if idx % 2 == 0 else "FFFFFF")
        set_cell_margins(c1, top=90, bottom=90, left=120, right=120)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        if idx == 3:
            r1.bold = True
            r1.font.color.rgb = COLOR_ACCENT
        else:
            r1.font.color.rgb = COLOR_TEXT_MAIN
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # --------------------------------------------------------------------------
    # 1. Resumen Ejecutivo
    # --------------------------------------------------------------------------
    add_header_styled(doc, "1. Resumen Ejecutivo", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Habiendo obtenido formalmente las credenciales y parámetros de conexión de la API Pública de Humand (v1), "
        "el presente informe detalla la arquitectura técnica, las etapas de ejecución y el plan riguroso de pruebas "
        "para interconectar de forma automatizada la Aplicación de Talento Humano (TH) con la plataforma Humand. "
        "El objetivo es gobernar el aprovisionamiento de empleados, el organigrama y las novedades laborales con alta "
        "disponibilidad, sin fricciones operativas y con absoluta fidelidad de datos."
    )
    
    add_callout_box(
        doc,
        "DISPONIBILIDAD DE CREDENCIALES",
        "Se confirma la recepción de la API Key corporativa para el entorno de producción Humand "
        "(https://api-prod.humand.co/public/api/v1). Se cuenta con la arquitectura de validación "
        "y los scripts de diagnóstico listos para dar inicio a la Fase 1 de pruebas.",
        border_color=COLOR_SECONDARY_HEX,
        bg_color=COLOR_BG_LIGHT_HEX
    )

    # --------------------------------------------------------------------------
    # 2. Principio Rector: Exclusión Total de Datos Salariales (DLP)
    # --------------------------------------------------------------------------
    add_header_styled(doc, "2. Principio Rector de Confidencialidad: Exclusión Absoluta de Bandas Salariales", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Por expresa instrucción de la Dirección y en estricto cumplimiento de las políticas de Prevención de Fuga de "
        "Datos (DLP) y confidencialidad estratégica de compensaciones, "
    )
    r_bold = p.add_run("QUEDA TERMINANTEMENTE PROHIBIDO TRANSMITIR, EXPONER O REGISTRAR CUALQUIER DATO SALARIAL EN HUMAND.")
    r_bold.bold = True
    r_bold.font.color.rgb = COLOR_ACCENT
    
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)
    p2.paragraph_format.line_spacing = 1.15
    p2.add_run(
        "Esta directiva establece un perímetro de seguridad infranqueable entre ambos sistemas:"
    )
    
    p_b1 = doc.add_paragraph(style='List Bullet')
    p_b1.paragraph_format.space_after = Pt(3)
    r = p_b1.add_run("Aislamiento Estricto en TH: ")
    r.bold = True
    p_b1.add_run("El módulo de Tabulador Salarial (bandas del 80% al 120%, medianas de referencia, percentiles, compa-ratio y salarios brutos/netos) residirá única y exclusivamente dentro de la base de datos interna de la Aplicación TH (PostgreSQL / InsForge).")

    p_b2 = doc.add_paragraph(style='List Bullet')
    p_b2.paragraph_format.space_after = Pt(3)
    r = p_b2.add_run("Filtrado en el Motor de Sincronización: ")
    r.bold = True
    p_b2.add_run("El adaptador de transformación descartará activamente cualquier atributo de compensación antes de construir los payloads JSON enviados a los endpoints de Humand.")

    p_b3 = doc.add_paragraph(style='List Bullet')
    p_b3.paragraph_format.space_after = Pt(6)
    r = p_b3.add_run("Omisión de Campos en Perfil: ")
    r.bold = True
    p_b3.add_run("No se creará ni vinculará ningún Profile Field dinámico relacionado a 'Banda Salarial', 'Nivel Salarial' o 'Sueldo' en la instancia corporativa de Humand.")

    add_callout_box(
        doc,
        "CONTROL DE SEGURIDAD DLP",
        "Humand se utilizará exclusivamente para comunicación institucional, organigrama, trámites de ausencias y "
        "experiencia del empleado. Las decisiones de mérito, compensaciones y equidad interna son competencia "
        "exclusiva de la Dirección y la Gerencia de TH a través de la Aplicación TH.",
        border_color=COLOR_ACCENT_HEX,
        bg_color="FFF5F5"
    )

    # --------------------------------------------------------------------------
    # 3. Matriz de Datos Autorizados: TH vs. Humand
    # --------------------------------------------------------------------------
    add_header_styled(doc, "3. Alcance de Datos Autorizados y Matriz de Sincronización", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "A continuación se especifica la clasificación explícita de los datos organizacionales, diferenciando "
        "los campos autorizados para sincronización de aquellos sujetos a estricta exclusión:"
    )
    
    t_map = doc.add_table(rows=11, cols=4)
    t_map.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_map.autofit = False
    set_table_borders(t_map, color="CBD5E0", sz="4")
    
    headers_map = ["Entidad TH", "Campo en TH", "Destino en Humand", "Estado / Regla de Gobernanza"]
    widths_map = [Inches(1.3), Inches(1.7), Inches(1.8), Inches(1.7)]
    
    # Encabezado de tabla
    hdr_row = t_map.rows[0]
    for i, h_text in enumerate(headers_map):
        cell = hdr_row.cells[i]
        cell.width = widths_map[i]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p_c = cell.paragraphs[0]
        r_c = p_c.add_run(h_text)
        r_c.font.name = "Calibri"
        r_c.font.size = Pt(8.5)
        r_c.bold = True
        r_c.font.color.rgb = RGBColor(255, 255, 255)
        
    data_map = [
        ("Empleado", "codigo_empleado", "employeeInternalId", "AUTORIZADO: Identificador unívoco"),
        ("Empleado", "documento_identidad", "Profile Field (Cédula)", "AUTORIZADO: Campo informativo"),
        ("Empleado", "nombres, apellidos", "firstName, lastName", "AUTORIZADO: Sanitizado sin redundancias"),
        ("Empleado", "email_corporativo", "email", "AUTORIZADO: Canal oficial de acceso"),
        ("Empleado", "telefono", "phoneNumber", "AUTORIZADO: Formato telefónico E.164"),
        ("Empleado", "fecha_ingreso", "hiringDate", "AUTORIZADO: Formato ISO YYYY-MM-DD"),
        ("Empleado", "di_supervisor", "relationships (BOSS)", "AUTORIZADO: Genera subordinados aut."),
        ("Empleado", "di_evaluador", "relationships (REVIEWER)", "AUTORIZADO: Aprobador de trámites"),
        ("Estructura", "departamento, cargo", "/departments, /job-positions", "AUTORIZADO: Catálogos maestros"),
        ("Compensación", "tabulador_id, salario", "NO APLICA / EXCLUIDO", "BLOQUEO TOTAL DLP: Confidencial TH")
    ]
    
    for idx, row_data in enumerate(data_map, start=1):
        row = t_map.rows[idx]
        bg = COLOR_BG_LIGHT_HEX if idx % 2 == 1 else "FFFFFF"
        if idx == 10:  # La fila de compensación
            bg = "FFF5F5"
            
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = widths_map[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            r_c = p_c.add_run(val)
            r_c.font.name = "Calibri"
            r_c.font.size = Pt(8)
            if idx == 10:
                r_c.bold = True
                r_c.font.color.rgb = COLOR_ACCENT
            else:
                r_c.font.color.rgb = COLOR_TEXT_MAIN

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --------------------------------------------------------------------------
    # 4. Reglas Técnicas Críticas de la API de Humand
    # --------------------------------------------------------------------------
    add_header_styled(doc, "4. Especificaciones Técnicas y Reglas de Oro de Humand", level=1)
    
    rules = [
        ("Header 'Basic' Obligatorio", "La API Key se envía mediante 'Authorization: Basic {api_key}'. El envío de 'Bearer' produce error 401 API_KEY_NOT_PROVIDED."),
        ("Sincronización Silenciosa con PUT /users", "El aprovisionamiento masivo se realiza con PUT /users. Este endpoint sobrescribe y NO envía correos automáticos al empleado, evitando el envío prematuro de cientos de correos de bienvenida (a diferencia de POST /users)."),
        ("Actualización Incremental con PATCH", "Para ajustes diarios (teléfono, correo, cambio de jefe), se emplea PATCH /users/{id}, modificando únicamente los campos transmitidos y preservando intactos los demás."),
        ("Esquemas de Paginación Diferenciados", "/users se gestiona con limit (máx. 50) y offset; los catálogos organizacionales (/departments, /job-positions) se consultan con page y limit."),
        ("Políticas de Rate Limiting y Resiliencia", "La API vigila la cuota mediante cabeceras RateLimit. En caso de HTTP 429, el motor ejecutará backoff exponencial respetando la cabecera Retry-After.")
    ]
    
    for title_r, desc_r in rules:
        p_r = doc.add_paragraph(style='List Bullet')
        p_r.paragraph_format.space_after = Pt(3)
        p_r.paragraph_format.line_spacing = 1.15
        r_t = p_r.add_run(f"{title_r}: ")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p_r.add_run(desc_r)
        r_d.font.name = "Calibri"
        r_d.font.color.rgb = COLOR_TEXT_MAIN

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --------------------------------------------------------------------------
    # 5. Fases del Plan de Implementación
    # --------------------------------------------------------------------------
    add_header_styled(doc, "5. Fases del Plan de Implementación", level=1)
    
    phases = [
        ("Fase 1: Diagnóstico y Sanity Check Inicial", 
         "Validación del apretón de manos (handshake) con la clave real contra /users/me, inspección de cabeceras de cuota y comprobación de permisos de lectura de la organización.",
         "Día 1"),
        ("Fase 2: Homologación de Catálogos Organizacionales", 
         "Extracción de Departamentos y Cargos desde TH y alta masiva en Humand (/departments/bulk y /job-positions/bulk). Creación de segmentaciones corporativas por Empresa y Sede.",
         "Días 2 - 3"),
        ("Fase 3: Construcción del Motor de Sincronización en TH", 
         "Despliegue del servicio seguro de transformación en el backend/Edge Functions. Implementación del filtro de exclusión de bandas salariales y cómputo de Hash SHA-256 para cambios incrementales.",
         "Días 4 - 6"),
        ("Fase 4: Desarrollo del Módulo de Integración en la Interfaz (UI)", 
         "Incorporación de una pantalla de administración en Aplicación TH con semáforo de conectividad, métricas de cuota, visualizador de discrepancias y botón de ejecución de Dry-Run.",
         "Días 7 - 8"),
        ("Fase 5: Aprovisionamiento Masivo Inicial (Batch Onboarding)", 
         "Carga inicial controlada en lotes de 20 empleados vía PUT /users. Verificación de creación del árbol jerárquico y relaciones BOSS/REVIEWER.",
         "Días 9 - 10"),
        ("Fase 6: Despliegue de Sincronización Incremental (Delta Sync)", 
         "Activación de sincronización periódica para altas, modificaciones vía PATCH y bajas formales (desactivación limpia con POST /deactivate y reasignación de subordinados).",
         "Días 11 - 12")
    ]
    
    t_fases = doc.add_table(rows=7, cols=3)
    t_fases.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_fases.autofit = False
    set_table_borders(t_fases, color="CBD5E0", sz="4")
    
    f_headers = ["Fase Operativa", "Alcance y Entregables Clave", "Duración Estimada"]
    f_widths = [Inches(1.8), Inches(3.7), Inches(1.0)]
    
    hdr_f = t_fases.rows[0]
    for i, h_text in enumerate(f_headers):
        cell = hdr_f.cells[i]
        cell.width = f_widths[i]
        set_cell_background(cell, COLOR_SECONDARY_HEX)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p_c = cell.paragraphs[0]
        r_c = p_c.add_run(h_text)
        r_c.font.name = "Calibri"
        r_c.font.size = Pt(8.5)
        r_c.bold = True
        r_c.font.color.rgb = RGBColor(255, 255, 255)
        
    for idx, (f_name, f_desc, f_time) in enumerate(phases, start=1):
        row = t_fases.rows[idx]
        bg = COLOR_BG_LIGHT_HEX if idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([f_name, f_desc, f_time]):
            cell = row.cells[c_idx]
            cell.width = f_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            r_c = p_c.add_run(val)
            r_c.font.name = "Calibri"
            r_c.font.size = Pt(8)
            if c_idx == 0:
                r_c.bold = True
                r_c.font.color.rgb = COLOR_PRIMARY
            elif c_idx == 2:
                r_c.bold = True
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                r_c.font.color.rgb = COLOR_TEXT_MAIN

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --------------------------------------------------------------------------
    # 6. Matriz de Pruebas de Calidad (Testing Strategy)
    # --------------------------------------------------------------------------
    add_header_styled(doc, "6. Matriz de Pruebas de Calidad y Criterios de Aceptación", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Para garantizar cero disrupciones y total integridad en producción, se ejecutará una batería de 7 pruebas "
        "secuenciales antes de habilitar la sincronización continua:"
    )
    
    tests = [
        ("T1", "Handshake & Identidad", "GET /users/me", "HTTP 200. Retorna nombre y rol de la cuenta de integración."),
        ("T2", "Auditoría de Rate Limiting", "Cabeceras HTTP", "Lectura correcta de cuotas y tiempo de refresco en segundos."),
        ("T3", "Lectura de Catálogos", "/departments y /job-positions", "Respuesta 200 con conteo de departamentos y puestos."),
        ("T4", "Piloto Unitario Ficticio", "PUT + PATCH + POST /deactivate", "Creación de 'EMP-TEST-PILOTO', actualización parcial y desactivación limpia con código 204."),
        ("T5", "Simulación Dry-Run", "Motor local en TH", "Reporte de conciliación 100% verificado sin realizar llamadas mutativas a Humand."),
        ("T6", "Aislamiento Salarial (DLP)", "Inspección de Payloads y Perfil", "Verificación estricta de ausencia total de datos de sueldo, banda o tabulador en Humand."),
        ("T7", "Batch Piloto Gradual", "Lote inicial de 10 empleados reales", "Aprovisionamiento exitoso sin envío de correos y con jerarquías enlazadas.")
    ]
    
    t_tests = doc.add_table(rows=8, cols=4)
    t_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_tests.autofit = False
    set_table_borders(t_tests, color="CBD5E0", sz="4")
    
    t_hdrs = ["ID", "Prueba de Calidad", "Método / Endpoint", "Criterio de Aprobación"]
    t_widths = [Inches(0.6), Inches(2.0), Inches(1.8), Inches(2.1)]
    
    hdr_t = t_tests.rows[0]
    for i, h_text in enumerate(t_hdrs):
        cell = hdr_t.cells[i]
        cell.width = t_widths[i]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p_c = cell.paragraphs[0]
        r_c = p_c.add_run(h_text)
        r_c.font.name = "Calibri"
        r_c.font.size = Pt(8.5)
        r_c.bold = True
        r_c.font.color.rgb = RGBColor(255, 255, 255)
        
    for idx, (t_id, t_name, t_meth, t_crit) in enumerate(tests, start=1):
        row = t_tests.rows[idx]
        bg = COLOR_BG_LIGHT_HEX if idx % 2 == 1 else "FFFFFF"
        if t_id == "T6":
            bg = "FFF5F5"
            
        for c_idx, val in enumerate([t_id, t_name, t_meth, t_crit]):
            cell = row.cells[c_idx]
            cell.width = t_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            r_c = p_c.add_run(val)
            r_c.font.name = "Calibri"
            r_c.font.size = Pt(8)
            if c_idx == 0:
                r_c.bold = True
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if t_id == "T6":
                    r_c.font.color.rgb = COLOR_ACCENT
                else:
                    r_c.font.color.rgb = COLOR_PRIMARY
            elif t_id == "T6":
                r_c.bold = True if c_idx == 1 else False
                r_c.font.color.rgb = COLOR_ACCENT
            else:
                r_c.font.color.rgb = COLOR_TEXT_MAIN

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --------------------------------------------------------------------------
    # 7. Gobierno de Credenciales y Mitigación de Riesgos
    # --------------------------------------------------------------------------
    add_header_styled(doc, "7. Seguridad de Credenciales y Plan de Contingencia", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Para blindar la operación, se adoptan los siguientes controles preventivos de ciberseguridad:"
    )
    
    sec_points = [
        ("Almacenamiento Seguro de la API Key", "La clave residirá exclusivamente en variables de entorno del servidor / Secrets de InsForge, jamás en el frontend ni en repositorios de código."),
        ("Control de Desactivación de Empleados (Offboarding)", "Cuando un empleado cambie a estado INACTIVO en TH, la API de Humand ejecutará POST /users/{id}/deactivate especificando motivo formal y reasignando automáticamente a sus subordinados a un nuevo supervisor para evitar orfandad en el organigrama."),
        ("Registro de Auditoría Transaccional", "Cada operación de sincronización almacenará en la base de datos interna la fecha/hora, el ID del empleado afectado, el hash SHA-256 del registro y el código HTTP de respuesta.")
    ]
    
    for s_t, s_d in sec_points:
        p_s = doc.add_paragraph(style='List Bullet')
        p_s.paragraph_format.space_after = Pt(3)
        p_s.paragraph_format.line_spacing = 1.15
        r_st = p_s.add_run(f"{s_t}: ")
        r_st.bold = True
        r_st.font.name = "Calibri"
        r_st.font.color.rgb = COLOR_PRIMARY
        r_sd = p_s.add_run(s_d)
        r_sd.font.name = "Calibri"
        r_sd.font.color.rgb = COLOR_TEXT_MAIN

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --------------------------------------------------------------------------
    # 8. Recomendación y Próximos Pasos Inmediatos
    # --------------------------------------------------------------------------
    add_header_styled(doc, "8. Conclusiones y Próximos Pasos Inmediatos", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "El presente plan garantiza una implementación estructurada, auditada y plenamente alineada con la "
        "política de cero exposición salarial hacia Humand. "
        "Se recomienda iniciar de inmediato con los siguientes hitos operativos:"
    )
    
    step1 = doc.add_paragraph(style='List Bullet')
    step1.paragraph_format.space_after = Pt(3)
    r1_b = step1.add_run("Paso Inmediato 1: ")
    r1_b.bold = True
    step1.add_run("Ejecución del script de diagnóstico ")
    r_code = step1.add_run("scripts/verify_humand_integration.py")
    r_code.font.name = "Consolas"
    step1.add_run(" ingresando la API Key en modo seguro para certificar la prueba T1 (Handshake) y T2 (Rate Limits).")

    step2 = doc.add_paragraph(style='List Bullet')
    step2.paragraph_format.space_after = Pt(3)
    r2_b = step2.add_run("Paso Inmediato 2: ")
    r2_b.bold = True
    step2.add_run("Construcción del componente UI de Sincronización en la Aplicación TH, dotando al equipo de Talento Humano del panel de supervisión y botón de simulación previa (Dry-Run).")

    step3 = doc.add_paragraph(style='List Bullet')
    step3.paragraph_format.space_after = Pt(14)
    r3_b = step3.add_run("Paso Inmediato 3: ")
    r3_b.bold = True
    step3.add_run("Sincronización inicial de Catálogos Maestros (Departamentos y Cargos) para preparar la estructura base de la organización en Humand.")

    # Bloque de Firmas y Responsables
    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False
    set_table_borders(sign_table, color="FFFFFF", sz="0")
    
    for c_i, (tit_cargo, area) in enumerate([("Líder de Sistemas y Arquitectura", "Dirección de Sistemas y Tecnología"), 
                                            ("Gerente Corporativo", "Gerencia de Talento Humano")]):
        c = sign_table.cell(0, c_i)
        c.width = Inches(3.2)
        set_cell_background(c, COLOR_BG_HEADER_HEX)
        set_cell_margins(c, top=140, bottom=140, left=140, right=140)
        p_sig = c.paragraphs[0]
        p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sig.paragraph_format.space_after = Pt(2)
        
        r_line = p_sig.add_run("_____________________________________\n")
        r_line.font.color.rgb = COLOR_TEXT_MUTED
        
        r_c = p_sig.add_run(f"{tit_cargo}\n")
        r_c.bold = True
        r_c.font.name = "Calibri"
        r_c.font.size = Pt(9.5)
        r_c.font.color.rgb = COLOR_PRIMARY
        
        r_a = p_sig.add_run(area)
        r_a.font.name = "Calibri"
        r_a.font.size = Pt(8.5)
        r_a.font.color.rgb = COLOR_TEXT_MUTED
        
    # Guardar documento
    output_dir = os.path.join(os.path.dirname(__file__), "..", "documentos")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "Informe Gerencial - Plan de Implementacion y Pruebas TH-Humand.docx")
    doc.save(output_path)
    print(f"Documento generado exitosamente en: {output_path}")

if __name__ == "__main__":
    build_document()
